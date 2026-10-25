#!/usr/bin/env python3
"""
Fill a .docx template and write a new document.

Text placeholders use {{name}}. Image placeholders use {{img_name}}
(underscore, so the name is valid in Jinja2).

Usage:
    generate.py <template_path> [options]
"""

import sys
import json
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from utils import (
    ensure_dependencies,
    scan_placeholders,
    format_placeholders_list,
    resolve_output_path,
    open_file,
    validate_template_path,
    validate_image_path,
    load_json_value,
    load_template_meta,
    trim_image_whitespace,
)


def generate_document(template_path, text_data, img_data, output_path, should_open=False):
    """Fill template placeholders and save the document."""
    from docxtpl import DocxTemplate, InlineImage
    from docx.shared import Cm, Mm

    template_path = validate_template_path(template_path)
    output_path = resolve_output_path(str(template_path), output_path)
    meta = load_template_meta(str(template_path))
    image_sizes = meta.get("images", {})

    context = {}
    placeholders = scan_placeholders(str(template_path))

    for name in placeholders["text"]:
        if name in text_data:
            context[name] = text_data[name]
        else:
            context[name] = f"[待填写: {name}]"

    tpl = DocxTemplate(str(template_path))
    for img_name in placeholders["img"]:
        context_key = f"img_{img_name}"
        if img_name in img_data:
            img_path = validate_image_path(img_data[img_name])
            if img_path:
                spec = image_sizes.get(img_name, {})
                width_cm = spec.get("width_cm")
                trimmed, px_w, px_h = trim_image_whitespace(img_path)
                kwargs = {}
                if width_cm:
                    w = float(width_cm)
                    kwargs["width"] = Cm(w)
                    if px_w > 0:
                        kwargs["height"] = Cm(w * (px_h / px_w))
                else:
                    kwargs["width"] = Mm(150)
                    if px_w > 0:
                        kwargs["height"] = Mm(150 * (px_h / px_w))
                context[context_key] = InlineImage(tpl, str(trimmed), **kwargs)
            else:
                context[context_key] = f"[待插入图片: {img_name}]"
        else:
            context[context_key] = f"[待插入图片: {img_name}]"

    tpl.render(context)
    tpl.save(str(output_path))
    print(f"Document written: {output_path}")

    if should_open:
        open_file(str(output_path))


def main():
    import argparse

    parser = argparse.ArgumentParser(description="Generate a Word document from a .docx template")
    parser.add_argument("template", help="Path to the .docx template file")
    parser.add_argument("--output", help="Output file path", default=None)
    parser.add_argument("--text-data", help="JSON string with text placeholder values", default=None)
    parser.add_argument("--text-data-file", help="UTF-8 JSON file with text placeholder values", default=None)
    parser.add_argument("--img-data", help="JSON string with image placeholder values (file paths)", default=None)
    parser.add_argument("--img-data-file", help="UTF-8 JSON file with image placeholder values (file paths)", default=None)
    parser.add_argument("--open", action="store_true", help="Open the generated file after creation")
    parser.add_argument("--scan-only", action="store_true", help="Only scan and display placeholders")

    args = parser.parse_args()
    ensure_dependencies()

    if args.scan_only:
        placeholders = scan_placeholders(args.template)
        print(format_placeholders_list(placeholders))
        return

    try:
        text_data = load_json_value(args.text_data, args.text_data_file)
        img_data = load_json_value(args.img_data, args.img_data_file)
    except json.JSONDecodeError as e:
        print(f"Error parsing JSON: {e}")
        sys.exit(1)

    generate_document(
        template_path=args.template,
        text_data=text_data,
        img_data=img_data,
        output_path=args.output,
        should_open=args.open,
    )


if __name__ == "__main__":
    main()
