#!/usr/bin/env python3
"""Rebuild the bundled fictional weekly-report template and its sample image."""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from utils import ensure_dependencies

ensure_dependencies()

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "assets" / "templates" / "weekly_report_template.docx"
META = ROOT / "assets" / "templates" / "weekly_report_template.meta.json"
PLACEHOLDERS = ROOT / "assets" / "templates" / "placeholders.txt"
SAMPLE_IMAGE = ROOT / "assets" / "examples" / "sample_figure.png"


def add_line(doc, text, size=12, bold=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return paragraph


def build_template():
    doc = Document()
    add_line(doc, "{{文档标题}}", size=18, bold=True)
    add_line(doc, "日期：{{日期}}")
    add_line(doc, "撰写人：{{撰写人}}")
    add_line(doc, "摘要", size=14, bold=True)
    add_line(doc, "{{摘要}}")
    add_line(doc, "本周进展", size=14, bold=True)
    add_line(doc, "{{本周进展}}")
    add_line(doc, "下周计划", size=14, bold=True)
    add_line(doc, "{{下周计划}}")
    add_line(doc, "附图", size=14, bold=True)
    add_line(doc, "{{img_附图}}")
    TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TEMPLATE)


def build_sample_image():
    image = Image.new("RGB", (800, 300), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 40, 760, 260), outline=(40, 90, 160), width=4)
    draw.text((70, 120), "SAMPLE FIGURE", fill=(40, 90, 160))
    SAMPLE_IMAGE.parent.mkdir(parents=True, exist_ok=True)
    image.save(SAMPLE_IMAGE)


def main():
    build_template()
    build_sample_image()
    META.write_text(
        '{\n  "images": {\n    "附图": {\n      "width_cm": 12\n    }\n  }\n}\n',
        encoding="utf-8",
    )
    PLACEHOLDERS.write_text(
        "\n".join(
            [
                "=== Text Placeholders ===",
                "  1. {{文档标题}}",
                "  2. {{日期}}",
                "  3. {{撰写人}}",
                "  4. {{摘要}}",
                "  5. {{本周进展}}",
                "  6. {{下周计划}}",
                "",
                "=== Image Placeholders ===",
                "  1. {{img_附图}}",
                "",
            ]
        ),
        encoding="utf-8",
    )
    print(TEMPLATE)
    print(SAMPLE_IMAGE)


if __name__ == "__main__":
    main()
