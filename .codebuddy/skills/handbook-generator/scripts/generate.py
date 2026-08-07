#!/usr/bin/env python3
"""
Handbook Generator - Generate Mode

Fills a .docx template with user-provided data and images.
Text placeholders use {{name}} syntax.
Image placeholders use {{img_name}} syntax (underscore, Jinja2-compatible).

Usage:
    generate.py <template_path> [options]

Options:
    --output <path>        Output file path (default: template dir + _generated suffix)
    --text-data <json>     JSON string with text placeholder values
    --img-data <json>      JSON string with image placeholder values (file paths)
    --open                 Open the generated file after creation
    --scan-only            Only scan and display placeholders, do not generate
"""

import sys
import json
from pathlib import Path

# Ensure utils is importable from the same directory
sys.path.insert(0, str(Path(__file__).parent))
from utils import (
    ensure_dependencies,
    scan_placeholders,
    format_placeholders_list,
    resolve_output_path,
    open_file,
    validate_template_path,
    validate_image_path,
)


def generate_handbook(template_path, text_data, img_data, output_path, should_open=False):
    """
    Generate a handbook by filling template placeholders with provided data.

    Args:
        template_path: Path to the .docx template file
        text_data: Dict mapping text placeholder names to their values
        img_data: Dict mapping image placeholder names to image file paths
        output_path: Path for the generated output file
        should_open: Whether to open the file after generation
    """
    from docxtpl import DocxTemplate, InlineImage
    from docx.shared import Mm

    template_path = validate_template_path(template_path)
    output_path = resolve_output_path(str(template_path), output_path)

    # Build the context dict for docxtpl
    context = {}

    # Scan template to discover all placeholders (for auto-filling missing ones)
    placeholders = scan_placeholders(str(template_path))

    # Fill text placeholders: use provided data, fallback to placeholder text for missing ones
    for name in placeholders['text']:
        if name in text_data:
            context[name] = text_data[name]
        else:
            context[name] = f"[待填写: {name}]"

    # Fill image placeholders: use provided data, fallback to placeholder text for missing ones
    tpl = DocxTemplate(str(template_path))
    for img_name in placeholders['img']:
        context_key = f"img_{img_name}"
        if img_name in img_data:
            img_path = validate_image_path(img_data[img_name])
            if img_path:
                context[context_key] = InlineImage(tpl, str(img_path), width=Mm(150))
            else:
                context[context_key] = f"[待插入图片: {img_name}]"
        else:
            context[context_key] = f"[待插入图片: {img_name}]"

    # Render the template
    tpl.render(context)

    # Save the generated file
    tpl.save(str(output_path))
    print(f"Handbook generated successfully: {output_path}")

    # Optionally open the file
    if should_open:
        open_file(str(output_path))


def main():
    import argparse

    parser = argparse.ArgumentParser(description='Generate a handbook from a .docx template')
    parser.add_argument('template', help='Path to the .docx template file')
    parser.add_argument('--output', help='Output file path', default=None)
    parser.add_argument('--text-data', help='JSON string with text placeholder values', default='{}')
    parser.add_argument('--img-data', help='JSON string with image placeholder values (file paths)', default='{}')
    parser.add_argument('--open', action='store_true', help='Open the generated file after creation')
    parser.add_argument('--scan-only', action='store_true', help='Only scan and display placeholders')

    args = parser.parse_args()

    # Ensure dependencies are available
    ensure_dependencies()

    # Scan-only mode: display placeholders and exit
    if args.scan_only:
        placeholders = scan_placeholders(args.template)
        print(format_placeholders_list(placeholders))
        return

    # Parse data from JSON strings
    try:
        text_data = json.loads(args.text_data)
    except json.JSONDecodeError as e:
        print(f"Error parsing text-data JSON: {e}")
        sys.exit(1)

    try:
        img_data = json.loads(args.img_data)
    except json.JSONDecodeError as e:
        print(f"Error parsing img-data JSON: {e}")
        sys.exit(1)

    # Generate the handbook
    generate_handbook(
        template_path=args.template,
        text_data=text_data,
        img_data=img_data,
        output_path=args.output,
        should_open=args.open
    )


if __name__ == '__main__':
    main()
